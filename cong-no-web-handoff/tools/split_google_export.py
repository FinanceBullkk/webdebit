"""Split a Google Sheets XLSX export into source-addressed JSONL tables.

The workbook remains the original export artifact. JSONL contains cached
display results from the XLSX export and is intended for read-only staging.
"""

from __future__ import annotations

import argparse
import hashlib
import json
from datetime import date, datetime, timezone
from pathlib import Path

import openpyxl


ROLES = {
    "Khách hàng": ("input", 6),
    "Số dư đầu": ("input", 6),
    "Chuyến hàng": ("input", 6),
    "Thu & cấn trừ": ("input", 6),
    "15 Vướng mắc 2026": ("exception", 3),
    "16 Thu 2026 khối cũ": ("support", 6),
    "09 Chuyến gốc": ("source_reference", 2),
    "10 Kỳ gốc": ("source_reference", 1),
    "11 Phí lẻ": ("source_reference", 1),
    "12 Cần kiểm tra": ("exception", 1),
    "Tổng quan": ("derived_view", 8),
    "Sổ khách": ("derived_view", 6),
    "06 Phiếu công nợ": ("derived_view", None),
    "17 Đối soát theo khách": ("historical_reconciliation", 6),
    "07 Đối soát": ("historical_reconciliation", 6),
    "08 Hướng dẫn": ("instructions", None),
}


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def encode(value):
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return int(value)
    return value


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("workbook", type=Path)
    parser.add_argument("metadata", type=Path)
    parser.add_argument("output_dir", type=Path)
    args = parser.parse_args()

    source = json.loads(args.metadata.read_text(encoding="utf-8"))
    args.output_dir.mkdir(parents=True, exist_ok=True)
    workbook = openpyxl.load_workbook(args.workbook, read_only=True, data_only=True)
    sheets = source["sheets"]
    if set(workbook.sheetnames) != {sheet["title"] for sheet in sheets}:
        raise ValueError("Workbook tabs do not match captured metadata")

    capture_id = args.output_dir.name
    manifest = {
        "capture_id": capture_id,
        "created_at_utc": datetime.now(timezone.utc).isoformat(),
        "source_spreadsheet_id": source["spreadsheet_id"],
        "source_title": source["title"],
        "source_url": source["url"],
        "source_modified_before_utc": source["modified_before_utc"],
        "source_modified_after_utc": source["modified_after_utc"],
        "source_modified_unchanged_during_export": source["modified_before_utc"] == source["modified_after_utc"],
        "capture_method": "Google Drive XLSX export; JSONL split from exported cached values",
        "consistency_limit": "Google export is not an atomic Sheet snapshot. Cached XLSX values and native Sheet formulas may differ; verify relevant live sentinels before financial claims.",
        "source_workbook": args.workbook.name,
        "source_workbook_sha256": sha256(args.workbook),
        "tabs": [],
    }

    for sheet in sheets:
        title = sheet["title"]
        gid = sheet["gid"]
        width = sheet["column_count"]
        role, header_row = ROLES[title]
        output = args.output_dir / f"gid-{gid}.jsonl"
        row_count = 0
        first_row = None
        last_row = None
        headers = None
        with output.open("w", encoding="utf-8", newline="\n") as stream:
            for source_row, row in enumerate(workbook[title].iter_rows(max_col=width, values_only=True), start=1):
                values = [encode(value) for value in row]
                if header_row == source_row:
                    headers = values
                if all(value is None or value == "" for value in values):
                    continue
                record = {
                    "capture_id": capture_id,
                    "source_file_id": source["spreadsheet_id"],
                    "sheet_gid": gid,
                    "source_row": source_row,
                    "values": values,
                }
                stream.write(json.dumps(record, ensure_ascii=False, separators=(",", ":")) + "\n")
                row_count += 1
                first_row = source_row if first_row is None else first_row
                last_row = source_row
        manifest["tabs"].append({
            "title": title,
            "gid": gid,
            "role": role,
            "header_row": header_row,
            "headers": headers,
            "grid_row_count": sheet["row_count"],
            "grid_column_count": width,
            "nonempty_export_rows": row_count,
            "first_nonempty_row": first_row,
            "last_nonempty_row": last_row,
            "jsonl_file": output.name,
            "jsonl_sha256": sha256(output),
        })

    target = args.output_dir / "manifest.json"
    target.write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({
        "capture_id": capture_id,
        "tabs": len(manifest["tabs"]),
        "rows": sum(tab["nonempty_export_rows"] for tab in manifest["tabs"]),
        "manifest": str(target),
    }, ensure_ascii=False))


if __name__ == "__main__":
    main()
